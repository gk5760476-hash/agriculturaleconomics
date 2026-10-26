"""Convert the raw source files in data/raw into tidy interim tables.

Raw files are never modified. Every function reads from data/raw and writes to
data/interim; derived indicators (growth, departures, shocks) belong in
transformations.py, not here.

Usage:
    python -I src/ingestion.py credit|rainfall|foodgrains|all
"""
from __future__ import annotations

import calendar
import json
import os
import re
import sys
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd

# ACREWS_ROOT lets the integration test run the pipeline on a scratch copy of the project.
ROOT = Path(os.environ.get("ACREWS_ROOT") or Path(__file__).resolve().parents[1])
RAW = ROOT / "data" / "raw"
INTERIM = ROOT / "data" / "interim"

FOCUS_STATES = ["BIHAR", "MAHARASHTRA"]

# --------------------------------------------------------------------------
# Credit: RBI DBIE, BSR-1 credit of SCBs by district / population group /
# bank group / occupation, annual (end-March), Rs crore.
# --------------------------------------------------------------------------
CREDIT_XLSX = RAW / "credit" / "rbi_dbie_scb_credit_district_occupation_annual.xlsx"
CREDIT_DATA_SHEETS = ["Data", "Data(1)", "Data(2)"]  # one table split at Excel's row limit
CREDIT_COLUMNS = [
    "year", "region", "state", "district", "population_group", "bank_group",
    "occupation_group", "occupation_subgroup", "n_accounts",
    "credit_limit_cr", "amount_outstanding_cr",
]


def _num(v):
    if v is None or v == "-" or v == "":
        return np.nan
    return float(v)


def ingest_credit() -> pd.DataFrame:
    from openpyxl import load_workbook

    wb = load_workbook(CREDIT_XLSX, read_only=True, data_only=True)
    rows = []
    for sheet in CREDIT_DATA_SHEETS:
        # Only the first sheet has a header; continuation sheets start straight with data.
        for r in wb[sheet].iter_rows(min_col=2, max_col=12, values_only=True):
            if r[0] is None or not str(r[0]).strip().isdigit():
                continue  # title, header, notes and blank rows
            rows.append((int(r[0]), *[None if x is None else str(x).strip() for x in r[1:8]],
                         _num(r[8]), _num(r[9]), _num(r[10])))
        print(f"  {sheet}: cumulative rows {len(rows):,}", flush=True)
    wb.close()

    df = pd.DataFrame(rows, columns=CREDIT_COLUMNS)
    for c in CREDIT_COLUMNS[1:8]:
        df[c] = df[c].astype("category")
    out = INTERIM / "credit"
    out.mkdir(parents=True, exist_ok=True)
    df.to_parquet(out / "bsr1_scb_credit_district_long.parquet", index=False)
    print(f"credit: {len(df):,} rows, years {df.year.min()}-{df.year.max()}")
    return df


# --------------------------------------------------------------------------
# Rainfall: IMD 0.25 deg gridded daily rainfall (mm). Years arrive either as
# NetCDF3 (RF25_ind<yyyy>_rfp25.nc) or as raw IMD binary (.grd, float32,
# laid out time x lat x lon). Both are reduced to region-month totals using
# masks built from the Survey of India district boundaries.
# --------------------------------------------------------------------------
LON = np.arange(66.5, 66.5 + 135 * 0.25, 0.25)
LAT = np.arange(6.5, 6.5 + 129 * 0.25, 0.25)
GEOJSON_ZIP = RAW / "geo" / "district_nwic_geojson.zip"


def _read_rain_year(year: int) -> np.ndarray:
    """Return daily rainfall (days, lat, lon) in mm with NaN outside India."""
    nc_path = RAW / "rainfall" / "imd_0p25_netcdf" / f"RF25_ind{year}_rfp25.nc"
    grd_path = RAW / "rainfall" / "imd_0p25_grd" / f"Rainfall_ind{year}_rfp25.grd"
    days = 366 if calendar.isleap(year) else 365
    if nc_path.exists():
        import netCDF4

        with netCDF4.Dataset(nc_path) as d:
            assert np.allclose(d["LONGITUDE"][:], LON) and np.allclose(d["LATITUDE"][:], LAT)
            arr = np.ma.filled(d["RAINFALL"][:].astype("float64"), np.nan)
    elif grd_path.exists():
        arr = np.fromfile(grd_path, dtype="<f4").astype("float64").reshape(days, len(LAT), len(LON))
        arr[arr <= -998] = np.nan
    else:
        raise FileNotFoundError(year)
    if arr.shape[0] != days:
        raise ValueError(f"{year}: expected {days} days, file has {arr.shape[0]}")
    return arr


def available_rain_years() -> list[int]:
    years = set()
    for p in (RAW / "rainfall").rglob("*_rfp25.*"):
        years.add(int(p.name.split("_ind")[1][:4]))
    return sorted(years)


def _load_districts():
    from pyproj import Transformer
    from shapely.geometry import shape
    from shapely.ops import transform

    with zipfile.ZipFile(GEOJSON_ZIP) as z:
        gj = json.loads(z.read("district_nwic.GeoJSON"))
    to_wgs84 = Transformer.from_crs("EPSG:7755", "EPSG:4326", always_xy=True).transform
    out = []
    for f in gj["features"]:
        p = f["properties"]
        out.append({
            "state_name": p["state_name"].strip(),
            "district": p["district"].strip(),
            "dtcode": p["dtcode"],
            "geometry": transform(to_wgs84, shape(f["geometry"])),
        })
    return out


def build_region_masks() -> pd.DataFrame:
    """Grid-cell membership: one row per (region, cell) using cell-centre-in-polygon.

    Regions: All-India (every valid IMD land cell), each focus state, and each
    district of the focus states. Districts too small to contain a cell centre
    fall back to the single nearest valid cell and are flagged.
    """
    import shapely

    valid = ~np.isnan(_read_rain_year(2011)[0])
    lon2d, lat2d = np.meshgrid(LON, LAT)
    pts = shapely.points(lon2d.ravel(), lat2d.ravel())
    valid_flat = valid.ravel()

    rows = [("All-India", "All-India", "", i, "imd_valid_cells") for i in np.flatnonzero(valid_flat)]
    districts = [d for d in _load_districts() if d["state_name"].upper() in FOCUS_STATES]
    for st in FOCUS_STATES:
        st_geom = shapely.union_all([d["geometry"] for d in districts if d["state_name"].upper() == st])
        idx = np.flatnonzero(shapely.contains_xy(st_geom, lon2d.ravel(), lat2d.ravel()) & valid_flat)
        rows += [(st.title(), st.title(), "", i, "cell_centre") for i in idx]
    for d in districts:
        g = d["geometry"]
        idx = np.flatnonzero(shapely.contains_xy(g, lon2d.ravel(), lat2d.ravel()) & valid_flat)
        method = "cell_centre"
        if len(idx) == 0:
            cand = np.flatnonzero(valid_flat)
            idx = [cand[np.argmin(shapely.distance(pts[cand], g.centroid))]]
            method = "nearest_cell_fallback"
        rows += [(f"{d['state_name'].title()}|{d['district']}", d["state_name"].title(), d["district"], i, method)
                 for i in idx]
    masks = pd.DataFrame(rows, columns=["region_id", "state", "district", "cell_index", "assignment"])
    masks["lat"] = lat2d.ravel()[masks.cell_index]
    masks["lon"] = lon2d.ravel()[masks.cell_index]
    out = INTERIM / "rainfall"
    out.mkdir(parents=True, exist_ok=True)
    masks.to_csv(out / "region_grid_masks.csv", index=False)
    return masks


def ingest_rainfall() -> pd.DataFrame:
    """Region x year x month rainfall: area-mean of the monthly cell totals (mm).

    Cells are weighted by cos(latitude) to approximate equal-area averaging.
    """
    masks = build_region_masks()
    cos_w = np.cos(np.deg2rad(np.repeat(LAT, len(LON))))
    groups = {rid: g.cell_index.to_numpy() for rid, g in masks.groupby("region_id")}
    meta = masks.drop_duplicates("region_id").set_index("region_id")[["state", "district"]]
    records = []
    for year in available_rain_years():
        daily = _read_rain_year(year)
        month = pd.date_range(f"{year}-01-01", periods=daily.shape[0]).month.to_numpy()
        for m in range(1, 13):
            tot = daily[month == m].sum(axis=0).ravel()  # NaN cells stay NaN
            for rid, idx in groups.items():
                v, w = tot[idx], cos_w[idx]
                ok = ~np.isnan(v)
                records.append((rid, year, m, float(np.sum(v[ok] * w[ok]) / np.sum(w[ok])) if ok.any() else np.nan,
                                int(ok.sum()), len(idx)))
        print(f"  rainfall {year} done", flush=True)
    df = pd.DataFrame(records, columns=["region_id", "year", "month", "rainfall_mm", "n_cells_valid", "n_cells"])
    df = df.join(meta, on="region_id")
    df["source_file_type"] = df.year.map(
        lambda y: "netcdf" if (RAW / "rainfall" / "imd_0p25_netcdf" / f"RF25_ind{y}_rfp25.nc").exists() else "grd")
    df = df[["region_id", "state", "district", "year", "month", "rainfall_mm", "n_cells_valid", "n_cells", "source_file_type"]]
    df.to_csv(INTERIM / "rainfall" / "imd_region_monthly_rainfall.csv", index=False)
    print(f"rainfall: {df.region_id.nunique()} regions, {df.year.nunique()} years")
    return df


# --------------------------------------------------------------------------
# Foodgrains: RBI Handbook of Statistics on Indian States 2024-25,
# Tables 62 (production, '000 t), 71 (area, '000 ha), 84 (yield, kg/ha).
# --------------------------------------------------------------------------
HBS_TABLES = {
    "states-hbs-t062.csv": ("foodgrain_production_kt", ": Production of Total Foodgrains"),
    "states-hbs-t071.csv": ("foodgrain_area_kha", ": Area of Total Foodgrains"),
    "states-hbs-t084.csv": ("foodgrain_yield_kg_ha", ": Estimates of Yield - Total Foodgrains"),
}


def ingest_foodgrains() -> pd.DataFrame:
    parts = []
    for fname, (var, suffix) in HBS_TABLES.items():
        t = pd.read_csv(RAW / "state_handbook" / fname)
        long = t.melt(id_vars="date", var_name="series", value_name="value")
        long["state"] = long.series.str.replace(suffix, "", regex=False).str.strip()
        long["variable"] = var
        long["source_table"] = fname
        parts.append(long.drop(columns="series"))
    df = pd.concat(parts, ignore_index=True)
    start = pd.to_datetime(df.pop("date")).dt.year
    df["crop_year"] = start.astype(str) + "-" + (start + 1).astype(str).str[-2:]
    df["crop_year_start"] = start
    wide = (df.pivot_table(index=["state", "crop_year", "crop_year_start"], columns="variable",
                           values="value", aggfunc="first").reset_index())
    wide.columns.name = None
    out = INTERIM / "foodgrains"
    out.mkdir(parents=True, exist_ok=True)
    wide.to_csv(out / "rbi_hbs_foodgrains_state_year.csv", index=False)
    print(f"foodgrains: {wide.state.nunique()} states, {wide.crop_year.min()}..{wide.crop_year.max()}")
    return wide


def ingest_handbook() -> pd.DataFrame:
    """Every RBI HBS state table in raw/state_handbook as one long table.

    `date` in these files is the start of the financial/agricultural year
    (2009-04-01 = 2009-10); rainfall Table 102 uses March dates. Values keep the
    published unit (credit Tables 153-158 are in rupees, not crore).
    """
    parts = []
    for p in sorted((RAW / "state_handbook").glob("states-hbs-t*.csv")):
        t = pd.read_csv(p)
        long = t.melt(id_vars="date", var_name="series", value_name="value").dropna(subset=["value"])
        split = long.series.str.split(": ", n=1, expand=True)
        long["state"], long["title"] = split[0].str.strip(), split[1].str.strip()
        long["table"] = p.stem.replace("states-hbs-", "")
        long["period_start"] = pd.to_datetime(long.date).dt.year
        parts.append(long[["table", "title", "state", "date", "period_start", "value"]])
    df = pd.concat(parts, ignore_index=True)
    out = INTERIM / "handbook"
    out.mkdir(parents=True, exist_ok=True)
    df.to_csv(out / "rbi_hbs_state_tables_long.csv", index=False)
    print(f"handbook: {df.table.nunique()} tables, {len(df):,} values")
    return df


# --------------------------------------------------------------------------
# Prices: WPI (Office of the Economic Adviser) financial-year indices for the
# three most recent bases, restricted to the crops in the design plan.
# --------------------------------------------------------------------------
WPI_FILES = {
    "2004-05": RAW / "prices" / "wpi" / "base_2004_05" / "yearlyf.xls",
    "2011-12": RAW / "prices" / "wpi" / "base_2011_12" / "fin_year_index.xls",
    "2022-23": RAW / "prices" / "wpi" / "base_2022_23" / "wpi_yearly_index_202606.xlsx",
}
# published name (lower-case, stripped) -> harmonised item
WPI_ITEMS = {
    "all commodities": "all_commodities", "i primary articles": "primary_articles",
    "rice": "paddy", "paddy": "paddy", "wheat": "wheat", "maize": "maize", "gram": "gram",
    "arhar": "arhar", "moong": "moong", "urad": "urad", "potato": "potato", "onion": "onion",
    "sugarcane": "sugarcane", "raw cotton": "raw_cotton", "soyabean": "soyabean",
    "rape & mustard seed": "rapeseed_mustard", "rapeseed & mustard seed": "rapeseed_mustard",
}


def ingest_wpi() -> pd.DataFrame:
    rows = []
    for base, path in WPI_FILES.items():
        x = pd.read_excel(path)
        name_col = next(c for c in x.columns if "name" in str(c).lower())
        names = x[name_col].astype(str).str.strip().str.lower().str.replace(r"\s+", " ", regex=True)
        # 2004-05 'Rice' is the paddy item (same code 1101010101 as later 'Paddy'); first match only,
        # so processed-food rows further down with the same name are not picked up.
        x = x.assign(item=names.map(WPI_ITEMS)).dropna(subset=["item"]).drop_duplicates("item")
        year_cols = [c for c in x.columns if re.search(r"\d{4}", str(c)) and c != name_col
                     and "code" not in str(c).lower() and "weight" not in str(c).lower()]
        for _, r in x.iterrows():
            rows.append((base, r["item"], r[name_col], 100.0, f"{base}"))  # base year = 100 by definition
            for c in year_cols:
                yrs = re.findall(r"\d{4}", str(c))
                fy = f"{yrs[0]}-{yrs[1][-2:]}" if len(yrs) == 2 else str(c)
                rows.append((base, r["item"], r[name_col], float(r[c]), fy))
    df = pd.DataFrame(rows, columns=["wpi_base", "item", "published_name", "index", "fiscal_year"])
    df = df.drop_duplicates(["wpi_base", "item", "fiscal_year"])
    out = INTERIM / "prices"
    out.mkdir(parents=True, exist_ok=True)
    df.to_csv(out / "wpi_items_fy_by_base.csv", index=False)
    print(f"wpi: {df.item.nunique()} items x {df.wpi_base.nunique()} bases")
    return df


if __name__ == "__main__":
    what = sys.argv[1] if len(sys.argv) > 1 else "all"
    for name, fn in [("foodgrains", ingest_foodgrains), ("handbook", ingest_handbook), ("wpi", ingest_wpi),
                     ("rainfall", ingest_rainfall), ("credit", ingest_credit)]:
        if what in (name, "all"):
            fn()
